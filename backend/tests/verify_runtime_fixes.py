import io
import requests
import openpyxl

BASE_URL = 'http://localhost:5173/api'

def run_verification():
    print('=====================================================')
    print('STARTING RUNTIME SYSTEM VERIFICATION ON LOCALHOST:5173')
    print('=====================================================')

    # 1. Student Login
    print('\n[1] Student Login via Vite proxy')
    s_login = requests.post(f'{BASE_URL}/auth/login', json={'email': '99240040955@klu.ac.in', 'password': '99240040955'}, timeout=10)
    assert s_login.status_code == 200, f'Student login failed: {s_login.text}'
    student_token = s_login.json()['token']
    student_headers = {'Authorization': f'Bearer {student_token}'}
    print('  -> Student logged in successfully.')

    # 2. Student lists published vivas
    print('\n[2] Student Vivas')
    vivas_res = requests.get(f'{BASE_URL}/student/vivas', headers=student_headers, timeout=10)
    assert vivas_res.status_code == 200, f'Get vivas failed: {vivas_res.text}'
    vivas = vivas_res.json().get('vivas', [])
    assert len(vivas) > 0, 'No vivas found'
    viva = vivas[0]
    print(f'  -> Found Viva ID {viva["id"]}: {viva["title"]}')

    # 3. Start Viva Session
    print('\n[3] Start Student Viva Session')
    start_res = requests.post(f'{BASE_URL}/student/vivas/{viva["id"]}/start', headers=student_headers, timeout=10)
    assert start_res.status_code == 200, f'Start viva failed: {start_res.text}'
    start_data = start_res.json()
    session_id = start_data['session']['id']
    q1 = start_data['current_question']
    print(f'  -> Session ID: {session_id}, Question 1 (Order {q1["question_order"]}): {q1["question_text"]}')

    def get_answers_for_question(q_text):
        q_lower = q_text.lower()
        if 'supervised' in q_lower:
            return (
                'Supervised learning uses labeled input-output pairs to learn an explicit mapping, whereas unsupervised learning uses unlabeled data to discover hidden patterns.',
                'Supervised learning works with labeled data to predict targets.'
            )
        elif 'cost' in q_lower:
            return (
                'The cost function like MSE quantifies squared prediction error, and its objective is to be minimized by adjusting weights via gradient descent.',
                'It measures the difference between predictions and actual targets.'
            )
        elif 'train-test' in q_lower or 'split' in q_lower:
            return (
                'Train-test split partitions data into training and validation sets to assess model generalization on unseen data and detect overfitting.',
                'It splits data to evaluate if the model works on unseen data.'
            )
        elif 'vanishing' in q_lower or 'relu' in q_lower:
            return (
                'Vanishing gradient happens when backpropagated gradients shrink exponentially through deep layers; ReLU maintains constant unit gradient for positive values.',
                'Gradients become too small in deep layers and ReLU stops saturation.'
            )
        else: # overfitting
            return (
                'Overfitting occurs when a model memorizes noise in training data; prevented by regularization techniques like Dropout and L1/L2 penalties.',
                'Model memorizes training noise instead of patterns, prevented by dropout.'
            )

    # 4. Submit Question 1 with a partially correct answer (Issue 1 verification)
    print('\n[4] Submit Question 1 (Partially Relevant Answer)')
    _, q1_partial = get_answers_for_question(q1['question_text'])
    print(f'  -> Student Answer: "{q1_partial}"')
    q1_sub = requests.post(
        f'{BASE_URL}/student/session-questions/{q1["id"]}/submit',
        headers=student_headers,
        json={'transcript': q1_partial},
        timeout=70,
    )
    assert q1_sub.status_code == 200, f'Q1 submit failed: {q1_sub.text}'
    q1_data = q1_sub.json()
    q1_eval = q1_data['evaluation']
    print(f'  -> Evaluated Score: {q1_eval["ai_score"]} / 10.0')
    print(f'  -> Evaluator Reason: {q1_eval["evaluation_reason"]}')
    print(f'  -> Model Improved Answer: {q1_eval.get("improved_answer", "N/A")[:80]}...')
    assert q1_eval['ai_score'] >= 5.0, f'Issue 1 Failure: Partially relevant answer scored too low ({q1_eval["ai_score"]})'
    print('  -> VERIFIED: Meaningfully partial answer awarded fair partial credit (>= 5.0)!')

    # 5. Submit Question 2 with a fully correct answer
    q2 = q1_data['next_question']
    print(f'\n[5] Submit Question 2 (Question {q2["question_order"]}): {q2["question_text"]}')
    q2_full, _ = get_answers_for_question(q2['question_text'])
    print(f'  -> Student Answer: "{q2_full}"')
    q2_sub = requests.post(
        f'{BASE_URL}/student/session-questions/{q2["id"]}/submit',
        headers=student_headers,
        json={'transcript': q2_full},
        timeout=70,
    )
    assert q2_sub.status_code == 200, f'Q2 submit failed: {q2_sub.text}'
    q2_data = q2_sub.json()
    q2_eval = q2_data['evaluation']
    print(f'  -> Evaluated Score: {q2_eval["ai_score"]} / 10.0')
    print(f'  -> Evaluator Reason: {q2_eval["evaluation_reason"]}')
    assert q2_eval['ai_score'] >= 8.0, f'Issue 1 Failure: Complete answer scored too low ({q2_eval["ai_score"]})'
    print('  -> VERIFIED: Accurate answer awarded full/high marks!')

    # 6. Submit Question 3 (Final Question)
    q3 = q2_data['next_question']
    print(f'\n[6] Submit Question 3 (Question {q3["question_order"]}): {q3["question_text"]}')
    q3_full, _ = get_answers_for_question(q3['question_text'])
    print(f'  -> Student Answer: "{q3_full}"')
    q3_sub = requests.post(
        f'{BASE_URL}/student/session-questions/{q3["id"]}/submit',
        headers=student_headers,
        json={'transcript': q3_full},
        timeout=70,
    )
    assert q3_sub.status_code == 200, f'Q3 submit failed: {q3_sub.text}'
    q3_data = q3_sub.json()
    assert q3_data.get('is_completed') is True, 'Session should be completed after Q3'
    result = q3_data.get('result', {})
    print(f'  -> Viva Completed! Final Score: {result.get("final_score")} / 20.0 (Best 2 of 3)')
    print(f'  -> Q1: {result.get("q1_score")}, Q2: {result.get("q2_score")}, Q3: {result.get("q3_score")}')
    print(f'  -> Counted Questions: #{result.get("counted_question_1")} and #{result.get("counted_question_2")}')

    # 7. Verify Post-Exam Report in Finalize Result (Issue 2 verification)
    print('\n[7] Verify Post-Exam Report Data (Issue 2)')
    reviews = result.get('questions_review', [])
    assert len(reviews) == 3, f'Expected 3 questions in questions_review, got {len(reviews)}'
    for r in reviews:
        print(f'  --- Question {r["question_order"]} ---')
        print(f'  Question: {r["question_text"]}')
        print(f'  Student Answer: {r["student_answer"]}')
        print(f'  Marks: {r["score"]} / {r["max_score"]} (Counted: {r["is_counted"]})')
        print(f'  Explanation: {r["reason"]}')
        print(f'  Improved Answer: {r["improved_answer"]}')
        assert len(r['question_text']) > 0
        assert len(r['student_answer']) > 0
        assert len(r['reason']) > 0
        assert len(r['improved_answer']) > 0
        assert r['score'] >= 0.0
    print('  -> VERIFIED: Post-exam questions_review contains all required fields!')

    # 8. Verify Student Results Page API
    print('\n[8] Verify Student Results History API (Issue 2)')
    res_list = requests.get(f'{BASE_URL}/student/results', headers=student_headers, timeout=10)
    assert res_list.status_code == 200, f'Get student results failed: {res_list.text}'
    results_data = res_list.json().get('results', [])
    assert len(results_data) > 0, 'No results found in student history'
    latest = results_data[0]
    assert 'questions_review' in latest, 'questions_review missing from student results history'
    assert len(latest['questions_review']) == 3
    print(f'  -> VERIFIED: Student results list provides post-exam review with {len(latest["questions_review"])} questions!')

    # 9. Faculty Login & Excel Marks Sheet Download (Issue 3 verification)
    print('\n[9] Faculty Login & Excel Download (Issue 3)')
    f_login = requests.post(f'{BASE_URL}/auth/login', json={'email': 'faculty@klu.ac.in', 'password': 'faculty123'}, timeout=10)
    assert f_login.status_code == 200, f'Faculty login failed: {f_login.text}'
    fac_token = f_login.json()['token']
    fac_headers = {'Authorization': f'Bearer {fac_token}'}
    print('  -> Faculty logged in successfully.')

    print('\n[10] Faculty Results API')
    f_results = requests.get(f'{BASE_URL}/faculty/vivas/{viva["id"]}/results', headers=fac_headers, timeout=10)
    assert f_results.status_code == 200, f'Faculty results failed: {f_results.text}'
    f_json = f_results.json()
    student_summaries = f_json.get('student_summaries', [])
    print(f'  -> Attended students in summary: {len(student_summaries)}')
    assert len(student_summaries) > 0, 'Expected at least 1 attended student summary'

    print('\n[11] Faculty Download Marks (.xlsx) via Vite Proxy')
    export_res = requests.get(
        f'{BASE_URL}/faculty/vivas/{viva["id"]}/results/export',
        headers=fac_headers,
        timeout=15,
    )
    assert export_res.status_code == 200, f'Excel export failed: {export_res.text}'
    ct = export_res.headers.get('Content-Type', '')
    cd = export_res.headers.get('Content-Disposition', '')
    print(f'  -> Response Status: {export_res.status_code}')
    print(f'  -> Content-Type: {ct}')
    print(f'  -> Content-Disposition: {cd}')
    assert 'openxmlformats' in ct or 'sheet' in ct or 'csv' in ct, f'Unexpected Content-Type: {ct}'
    assert 'attachment' in cd, f'Unexpected Content-Disposition: {cd}'

    wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
    assert 'Student Marks' in wb.sheetnames
    ws = wb.active
    print(f'  -> Excel Title Banner: {ws["A1"].value}')
    print(f'  -> Excel Sub-banner: {ws["A2"].value}')
    headers = [cell.value for cell in ws[4]]
    print(f'  -> Headers (Row 4): {headers}')
    student_row = [cell.value for cell in ws[5]]
    print(f'  -> Student Row (Row 5): {student_row}')
    assert student_row[0] == viva['title']
    assert student_row[7] == result.get('final_score')
    print('  -> VERIFIED: Excel (.xlsx) generated and validated with real student evaluated marks!')

    print('\n=====================================================')
    print('ALL 3 RUNTIME ISSUES VERIFIED AND FIXED SUCCESSFULLY!')
    print('=====================================================')

if __name__ == '__main__':
    run_verification()
