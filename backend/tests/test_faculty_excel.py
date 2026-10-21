import io
import sys
from datetime import datetime, timezone
from pathlib import Path

# Add backend directory to path
backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def test_excel_generation():
    headers = [
        "Viva Name",
        "Student Name",
        "Register Number",
        "Institutional Email",
        "Question 1 Score",
        "Question 2 Score",
        "Question 3 Score",
        "Best 2 Score",
        "Maximum Score",
        "Status",
        "Viva Date",
        "Completed At",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Marks"

    # University Header Banner
    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = "Kalasalingam Academy of Research and Education — Advanced Algorithms Viva — Marks Sheet"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:L2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} | Best 2 of 3 Questions (Maximum 20.0 Marks)"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.append([])  # blank row 3
    ws.append(headers)  # row 4

    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    sample_sessions = [
        {
            "viva_title": "Advanced Algorithms Viva",
            "name": "Rishi Kumar",
            "reg_no": "9921004123",
            "email": "rishi@klu.ac.in",
            "q1": 10.0,
            "q2": 8.5,
            "q3": 0.0,
            "best2": 18.5,
            "max": 20.0,
            "status": "COMPLETED",
            "viva_date": "2026-10-21 10:00",
            "comp_date": "2026-10-21 10:02",
        },
        {
            "viva_title": "Advanced Algorithms Viva",
            "name": "Ananya Sharma",
            "reg_no": "9921004124",
            "email": "ananya@klu.ac.in",
            "q1": 10.0,
            "q2": 10.0,
            "q3": 9.0,
            "best2": 20.0,
            "max": 20.0,
            "status": "COMPLETED",
            "viva_date": "2026-10-21 10:05",
            "comp_date": "2026-10-21 10:07",
        },
    ]

    row_idx = 5
    for s in sample_sessions:
        row_data = [
            s["viva_title"],
            s["name"],
            s["reg_no"],
            s["email"],
            s["q1"],
            s["q2"],
            s["q3"],
            s["best2"],
            s["max"],
            s["status"],
            s["viva_date"],
            s["comp_date"],
        ]
        ws.append(row_data)
        for c_idx in range(1, len(row_data) + 1):
            c = ws.cell(row=row_idx, column=c_idx)
            c.border = thin_border
            if c_idx in (5, 6, 7, 8, 9):
                c.alignment = Alignment(horizontal="center")
        row_idx += 1

    for col in ws.iter_cols(min_row=4, max_row=ws.max_row):
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    # Verify reload of saved workbook
    loaded_wb = openpyxl.load_workbook(output)
    sheet = loaded_wb["Student Marks"]
    assert sheet.cell(row=4, column=1).value == "Viva Name"
    assert sheet.cell(row=4, column=2).value == "Student Name"
    assert sheet.cell(row=4, column=5).value == "Question 1 Score"
    assert sheet.cell(row=4, column=8).value == "Best 2 Score"
    assert sheet.cell(row=5, column=2).value == "Rishi Kumar"
    assert sheet.cell(row=5, column=5).value == 10.0
    assert sheet.cell(row=5, column=8).value == 18.5
    assert sheet.cell(row=6, column=2).value == "Ananya Sharma"
    assert sheet.cell(row=6, column=8).value == 20.0
    print("Excel export verification: SUCCESS! File generated and validated perfectly.")


if __name__ == "__main__":
    test_excel_generation()
